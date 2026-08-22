# -*- coding: utf-8 -*-
"""只读验证：测试2 格式解析 + 推导（月租/年租/税率/税后）"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from datetime import datetime

path = r"F:\DieJiaTai\项目开发\水库管理平台\项目相关文件\聊天文件\场地信息关联表（测试2）.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
ws = next((w for w in wb.worksheets if "合同" in w.title or "收付" in w.title), wb[wb.sheetnames[0]])
rows = list(ws.iter_rows(min_row=1, values_only=True))

def _parse_date(val):
    if val is None: return None
    if isinstance(val, datetime): return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if not s: return None
    for fmt in ("%Y.%m.%d", "%Y/%m/%d", "%Y-%m-%d", "%Y年%m月%d日"):
        try: return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError: continue
    return None
def _num(val):
    if val is None: return None
    try: return float(val)
    except (ValueError, TypeError): return None
def _str(val):
    if val is None: return None
    s = str(val).strip()
    return s if s else None

current_landlord = None
for row_idx, row in enumerate(rows[3:], start=4):
    ct_venue = _str(row[0]); landlord_name = _str(row[1]); region = _str(row[2]); address = _str(row[3])
    rent_amount = _num(row[4]); monthly_rent = _num(row[5]); elec_pay = _num(row[6]); pm_v = _str(row[7])
    deposit_v = _num(row[8]); adv = _num(row[9]); sd_v = _parse_date(row[10]); ed_v = _parse_date(row[11])
    ct_brand = _str(row[13]); brand = _str(row[14]); rcm = _str(row[15]); unit = _num(row[16]); cabs = _num(row[17])
    bmr = _num(row[18]); brent = _num(row[19]); bep = _num(row[21]); pm_b = _str(row[22])
    first = _num(row[24]); elec_tax_pct = _num(row[26]); rent_tax_pct = _num(row[27])
    sd_b = _parse_date(row[28]); ed_b = _parse_date(row[29])

    if ct_venue and "场地" in ct_venue:
        if monthly_rent is None and rent_amount:
            monthly_rent = round(rent_amount / 12, 2)
        current_landlord = landlord_name
        print(f"[{row_idx}] 场地合同: 场地={landlord_name} 区域={region} 地址={address} 年租={rent_amount} 月租(推导)={monthly_rent} 电费单价={elec_pay} 付款={pm_v} 押金={deposit_v} 提前={adv} 起止={sd_v}~{ed_v}")
    if ct_brand and "品牌" in ct_brand:
        is_fixed = rcm and "固定" in rcm
        rcm_out = "固定价格" if is_fixed else "按柜子数量"
        if bmr is None:
            bmr = round((unit or 0) / 12, 2) if is_fixed and unit else (round(unit * cabs, 2) if unit and cabs else None)
        if brent is None:
            brent = unit if is_fixed else (round(bmr * 12, 2) if bmr else None)
        etr = round((elec_tax_pct or 0) / 100, 4) if elec_tax_pct else None
        rtr = round((rent_tax_pct or 0) / 100, 4) if rent_tax_pct else None
        pte = round(bep * (1 - etr), 4) if bep and etr else None
        ptr = round(bmr * (1 - rtr), 2) if bmr and rtr else None
        print(f"  → 品牌合同: 关联场地={current_landlord} 品牌={brand} 计费={rcm_out} 单柜={unit} 柜数={cabs} 月租(推导)={bmr} 年租(推导)={brent} 电费收款={bep} 收款方式={pm_b} 首月={first}")
        print(f"     电费税率={etr} 租金税率={rtr} 税后电费单价={pte} 税后月租={ptr} 起止={sd_b}~{ed_b}")
