# -*- coding: utf-8 -*-
"""调试：打印测试2 第9行（铁塔）原始单元格值"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl

path = r"F:\DieJiaTai\项目开发\水库管理平台\项目相关文件\聊天文件\场地信息关联表（测试2）.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
ws = next((w for w in wb.worksheets if "合同" in w.title or "收付" in w.title), wb[wb.sheetnames[0]])
rows = list(ws.iter_rows(min_row=1, values_only=True))

cols = "ABCDEFGHIJKLMNOPQRSTUVWXYZAAABACADAEAF"
for idx, row in enumerate(rows[3:], start=4):
    if row[14] and "铁塔" in str(row[14]):
        print(f"第{idx}行 铁塔:")
        for i, v in enumerate(row):
            if v is not None:
                col = cols[i] if i < 26 else f"A{cols[i-26]}"
                print(f"  {col}({i}): {v!r}")
