# -*- coding: utf-8 -*-
"""验证：同一行同时是 场地合同 + 品牌合同 时，两者都处理（elif→if 修复）"""
import openpyxl
from datetime import datetime

path = r"F:\DieJiaTai\项目开发\水库管理平台\项目相关文件\聊天文件\场地信息关联表（测试2）.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
ws = next((w for w in wb.worksheets if "合同" in w.title or "收付" in w.title), wb[wb.sheetnames[0]])
rows = list(ws.iter_rows(min_row=1, values_only=True))

def _str(v):
    if v is None: return None
    s = str(v).strip()
    return s if s else None

print("== 各行识别结果（is_venue / is_brand / 场地上下文）==")
current_landlord = None
for row_idx, row in enumerate(rows[3:], start=4):
    ct_venue = _str(row[0])
    ct_brand = _str(row[13])
    landlord_name = _str(row[1])
    brand_name = _str(row[14])
    is_venue = ct_venue and "场地" in ct_venue
    is_brand = ct_brand and "品牌" in ct_brand
    actions = []
    if is_venue:
        current_landlord = landlord_name
        actions.append(f"处理场地合同({landlord_name})")
    if is_brand:
        actions.append(f"处理品牌合同({brand_name} → 关联场地 {current_landlord})")
    print(f"  [第{row_idx}行] {' 且 '.join(actions) if actions else '跳过(合并区非首行或无数据)'}")
